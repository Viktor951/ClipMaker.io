from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_border(cell, **kwargs):
    """
    Set cell`s border
    Usage:
    set_cell_border(
        cell,
        top={"sz": 12, "val": "single", "color": "#FF0000", "space": "0"},
        bottom={"sz": 12, "color": "#00FF00", "val": "single"},
        start={"sz": 24, "val": "dashed", "shadow": "true"},
        end={"sz": 12, "val": "dashed"},
    )
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'start', 'bottom', 'end', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

def add_styled_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, text in enumerate(row_data):
            row_cells[i].text = str(text)
            
    # Set column widths roughly
    if len(headers) == 4:
        widths = [0.5, 1.5, 3.0, 1.5]
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Pt(width * 72)
                
    return table

doc = Document()

# Title
title = doc.add_heading('ClipMakerAI — Refatoração Profunda & Estabilização (Produção)', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Missão: Estancar vazamentos (Memory/File Leaks), separar responsabilidades (Service Pattern) e otimizar concorrência de hardware (NVENC vs Whisper na VRAM).')

doc.add_heading('Diagnóstico e Execução Completa', level=1)
doc.add_paragraph('Após testes rigorosos de End-to-End (E2E), executei uma reestruturação arquitetural completa. Abaixo estão detalhadas todas as modificações aplicadas ao projeto.')

# Table 1
doc.add_heading('🔴 Vazamentos Críticos Resolvidos (Leaks)', level=2)
headers1 = ['#', 'Arquivo', 'Bug Resolvido / Modificação', 'Impacto']
rows1 = [
    ['1', '🐍 main.py', 'O vídeo de origem .mp4 nunca era deletado do disco se a rota executasse com sucesso. Adicionado finally estrito.', 'Evita lotar o HD do servidor rapidamente (Storage Exhaustion)'],
    ['2', '🐍 engine.py', 'Arquivos .srt ficavam órfãos e não eram apagados caso o FFmpeg sofresse um crash. Adicionado try/finally no render.', 'Previne o acúmulo de arquivos residuais invisíveis'],
    ['3', '🐍 ai_service.py', 'O modelo Whisper ficava na VRAM após uso, concorrendo com o FFmpeg. Aplicado del model e torch.cuda.empty_cache().', 'Evita crashes de "Out of Memory" (OOM) na RTX 3050']
]
add_styled_table(doc, headers1, rows1)

# Table 2
doc.add_heading('🟢 Refatoração Arquitetural (Nova Camada de Serviços)', level=2)
doc.add_paragraph('O monolítico engine.py de ~420 linhas foi quebrado em microsserviços internos especializados para adoção de padrões corporativos.')

headers2 = ['#', 'Arquivo', 'Refatoração', 'Impacto']
rows2 = [
    ['4', '🐍 engine.py', 'Reescrito completamente. Agora atua apenas como Orquestrador das chamadas, caindo para ~70 linhas de código limpo.', 'Escalabilidade, legibilidade e fácil manutenção futura'],
    ['5', '🐍 ai_service.py', '[NEW] Isola a inteligência artificial (Faster Whisper) e a heurística matemática de cortes de engajamento.', 'Separação clara do consumo pesado de IA'],
    ['6', '🐍 video_service.py', '[NEW] Isola o Face Tracking (OpenCV) e comandos FFmpeg (NVENC com fallback automático para CPU).', 'Desacopla o processamento de mídia do resto do código'],
    ['7', '🐍 db_service.py', '[NEW] Configuração segura do Prisma ORM (Prepared Statements) e consultas Anti-N+1 via Eager Loading (include).', 'Consultas rápidas e protegidas nativamente contra SQL Injection']
]
add_styled_table(doc, headers2, rows2)

# Table 3
doc.add_heading('🟡 Estabilização de Frontend & Experiência (UI)', level=2)
headers3 = ['#', 'Arquivo', 'Modificação', 'Benefício']
rows3 = [
    ['8', '🟨 app.js', 'Remoção de chamadas console.log residuais vazadas na versão final de produção.', 'Segurança da informação e console mais limpo'],
    ['9', '🟨 app.js', 'Remoção completa do bloqueante alert() em blocos catch.', 'Fim de travamentos bruscos na tela do usuário'],
    ['10', '🟨 app.js', 'Injeção dinâmica no DOM (error-banner) com temporizador de 8s para lidar com timeouts da API de forma suave.', 'Tratamento visual elegante, similar a plataformas SaaS']
]
add_styled_table(doc, headers3, rows3)

doc.add_heading('🚀 Suíte de Testes Adicionada', level=2)
doc.add_paragraph('[NEW] 🧪 test_pipeline.py\nCriado na raiz do projeto, este script utiliza o TestClient do FastAPI para injetar dinamicamente um vídeo sintético (gerado por FFmpeg), testando as rotas de ponta a ponta sem necessidade de subir as portas de rede. O script rastreia e aprova a limpeza absoluta da pasta /temp_videos/ logo após a execução.')

doc.save('Relatorio_Refatoracao_ClipMakerAI.docx')
print("Document generated successfully.")
